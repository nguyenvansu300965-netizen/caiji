from openpyxl import load_workbook

from src.services import exporter


def test_exports_filtered_rows_to_xlsx(tmp_path, monkeypatch):
    row = [
        "Global Parts",
        "+4930123456",
        "DE",
        "verified",
        "wholesaler",
        "medium",
        "https://example.com",
        "https://example.com/company",
        "2026-01-02 03:04:05",
    ]
    monkeypatch.setattr(exporter, "_iter_export_rows", lambda _filters: iter([row]))
    output = tmp_path / "leads.xlsx"
    count = exporter.export_companies(str(output), {"country": "DE"})
    workbook = load_workbook(output)
    sheet = workbook.active
    assert count == 1
    assert sheet["A2"].value == "Global Parts"
    assert sheet["B2"].value == "+4930123456"


def test_excel_export_removes_illegal_control_characters(tmp_path, monkeypatch):
    row = [
        "i-Lab\x0b Stretching & Fitness",
        "+85221234567",
        "HK",
        "verified",
        "fitness",
        "low",
        "",
        "https://example.com",
        "2026-01-02 03:04:05",
    ]
    monkeypatch.setattr(exporter, "_iter_export_rows", lambda _filters: iter([row]))
    output = tmp_path / "safe.xlsx"
    exporter.export_companies(str(output), {})
    sheet = load_workbook(output).active
    assert sheet["A2"].value == "i-Lab Stretching & Fitness"
